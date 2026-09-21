from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Protocol

import httpx
from bs4 import BeautifulSoup
from pydantic import ValidationError

from .models import CompanyRecord, EvaluationResult, PipelineRow, RowStatus, ScrapedContent


class Provider(Protocol):
    def evaluate(self, *, company: CompanyRecord, content: ScrapedContent, prompt: str) -> EvaluationResult:
        ...


class ContentStore(Protocol):
    def get(self, url: str) -> ScrapedContent | None: ...

    def put(self, content: ScrapedContent) -> None: ...


def load_companies(path: str | Path) -> list[CompanyRecord]:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        return [CompanyRecord.model_validate(row) for row in rows]

    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        header = [str(v or "").strip() for v in next(iterator)]
        return [
            CompanyRecord.model_validate(
                {header[i]: (values[i] if i < len(values) and values[i] is not None else "") for i in range(len(header))}
            )
            for values in iterator
        ]

    raise ValueError("input must be CSV or XLSX")


def load_prompt(project: str, root: str | Path = "prompts") -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}", project):
        raise ValueError("unsafe project name")
    path = Path(root) / f"{project}.txt"
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        raise ValueError("prompt is empty")
    return text


class RawContentStore:
    def __init__(self, root: str | Path = "data/raw") -> None:
        self.root = Path(root)

    @staticmethod
    def key(url: str) -> str:
        normalized = url.strip().lower().rstrip("/")
        return hashlib.sha256(normalized.encode()).hexdigest()[:24]

    def path_for(self, url: str) -> Path:
        return self.root / f"{self.key(url)}.json"

    def get(self, url: str) -> ScrapedContent | None:
        path = self.path_for(url)
        if not path.exists():
            return None
        return ScrapedContent.model_validate_json(path.read_text(encoding="utf-8"))

    def put(self, content: ScrapedContent) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        path = self.path_for(content.website)
        path.write_text(content.model_dump_json(indent=2), encoding="utf-8")


class WebsiteFetcher:
    def __init__(self, timeout_seconds: float = 20.0) -> None:
        self.timeout_seconds = timeout_seconds

    def fetch(self, company: CompanyRecord) -> ScrapedContent:
        with httpx.Client(
            timeout=self.timeout_seconds,
            follow_redirects=True,
            headers={"User-Agent": "AICompanyEvaluationPipeline/0.1 (+portfolio-demo)"},
        ) as client:
            response = client.get(str(company.website))
            response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")
        for tag in soup(["script", "style", "noscript", "svg", "template"]):
            tag.decompose()

        text = re.sub(r"\s+", " ", " ".join(soup.stripped_strings)).strip()
        if not text:
            raise ValueError("no website text extracted")

        digest = hashlib.sha256(text.encode()).hexdigest()
        return ScrapedContent(
            company_name=company.company_name,
            website=str(company.website),
            combined_text=text,
            content_sha256=digest,
        )


def run_pipeline(
    companies: list[CompanyRecord],
    *,
    project: str,
    prompt: str,
    provider: Provider,
    store: ContentStore,
    fetcher: WebsiteFetcher | None = None,
) -> list[PipelineRow]:
    rows: list[PipelineRow] = []

    for company in companies:
        cached = store.get(str(company.website))
        cache_hit = cached is not None

        if cached is None:
            if fetcher is None:
                rows.append(PipelineRow(
                    company_name=company.company_name,
                    website=str(company.website),
                    project=project,
                    status=RowStatus.CACHE_MISS,
                    error="no cached content",
                ))
                continue

            try:
                cached = fetcher.fetch(company)
                store.put(cached)
            except Exception:
                rows.append(PipelineRow(
                    company_name=company.company_name,
                    website=str(company.website),
                    project=project,
                    status=RowStatus.SCRAPE_ERROR,
                    error="website content could not be fetched",
                ))
                continue

        try:
            evaluation = provider.evaluate(company=company, content=cached, prompt=prompt)
            evaluation = EvaluationResult.model_validate(evaluation)
            rows.append(PipelineRow(
                company_name=company.company_name,
                website=str(company.website),
                project=project,
                status=RowStatus.SUCCESS,
                score=evaluation.score,
                recommendation=evaluation.recommendation,
                rationale=evaluation.rationale,
                cache_hit=cache_hit,
            ))
        except Exception as exc:
            rows.append(PipelineRow(
                company_name=company.company_name,
                website=str(company.website),
                project=project,
                status=RowStatus.EVALUATION_ERROR,
                cache_hit=cache_hit,
                error="provider output failed validation" if isinstance(exc, ValidationError) else "provider evaluation failed",
            ))

    return rows


def write_results(rows: list[PipelineRow], output_dir: str | Path, project: str) -> tuple[Path, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = [row.model_dump(mode="json") for row in rows]

    json_path = output_dir / f"{project}_results.json"
    csv_path = output_dir / f"{project}_results.csv"

    json_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    fields = ["company_name", "website", "project", "status", "score", "recommendation", "rationale", "cache_hit", "error"]
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows({k: item.get(k) for k in fields} for item in payload)

    return csv_path, json_path
